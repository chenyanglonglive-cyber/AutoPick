from __future__ import annotations

import hashlib
import json
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from openpyxl.utils.units import pixels_to_EMU


RESOURCE_ROOT = Path(__file__).resolve().parents[2] / "resources" / "checklists"
QUALITY_RESOURCE = RESOURCE_ROOT / "quality_v1" / "Quality list.xlsx"
QUALITY_SHA256 = "EC05C7B51702527E519AC8A112CC51A05D69ED95CF194637D5F7ED27D104C8BF"


@dataclass(frozen=True)
class ChecklistSlotSpec:
    item_key: str
    ordinal: int
    section: str
    label: str
    sheet_name: str
    label_cell: str
    image_cell: str


@dataclass(frozen=True)
class ChecklistType:
    id: str
    label: str
    version: str
    resource_path: Path
    sha256: str
    slots: tuple[ChecklistSlotSpec, ...]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _merged_anchor(ws, row: int) -> bool:
    return any(r.min_row == row and r.max_row == row and r.min_col == 1 and r.max_col >= 2 for r in ws.merged_cells.ranges)


def _build_quality_slots(path: Path) -> tuple[ChecklistSlotSpec, ...]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    ws = workbook["Photo report Tool"]
    slots: list[ChecklistSlotSpec] = []
    section = ""
    occurrences: dict[str, int] = {}
    for row in range(1, ws.max_row + 1):
        a = ws.cell(row, 1)
        b = ws.cell(row, 2)
        if _merged_anchor(ws, row):
            section = str(a.value or "").strip()
            continue
        # Onsite markers are instructional separators, not photo fields.
        values = [(a, a.value), (b, b.value)]
        for cell, value in values:
            if value is None or cell.style_id not in {5, 9}:
                continue
            label = str(value).strip()
            if not label:
                continue
            if label.lower() == "onsite:":
                continue
            image_row = cell.row - 1
            image_cell = f"{get_column_letter(cell.column)}{image_row}"
            if ws[image_cell].value is not None or ws.row_dimensions[image_row].height is None:
                raise ValueError(f"固定清单图片位置校验失败: {ws.title}!{image_cell}")
            occurrences[label] = occurrences.get(label, 0) + 1
            item_key = f"quality_v1/{section or 'default'}/{cell.coordinate}/{occurrences[label]}"
            slots.append(ChecklistSlotSpec(item_key, len(slots) + 1, section, label, ws.title, cell.coordinate, image_cell))
    if len(slots) != 199:
        raise ValueError(f"固定质量清单应有 199 个图片字段，实际发现 {len(slots)} 个")
    return tuple(slots)


def checklist_types() -> list[ChecklistType]:
    if not QUALITY_RESOURCE.is_file():
        raise FileNotFoundError(f"内置清单资源不存在: {QUALITY_RESOURCE}")
    digest = _sha256(QUALITY_RESOURCE)
    if digest != QUALITY_SHA256:
        raise ValueError("Quality list.xlsx 校验失败，可能被意外修改")
    return [ChecklistType("quality_v1", "Quality list（质量审核）", "1", QUALITY_RESOURCE, digest, _build_quality_slots(QUALITY_RESOURCE))]


def get_checklist_type(type_id: str) -> ChecklistType:
    for checklist in checklist_types():
        if checklist.id == type_id:
            return checklist
    raise ValueError(f"暂不支持清单类型: {type_id}")


def slot_dicts(type_id: str = "quality_v1") -> list[dict]:
    checklist = get_checklist_type(type_id)
    return [slot.__dict__.copy() for slot in checklist.slots]


def copy_template(type_id: str, destination: Path) -> None:
    checklist = get_checklist_type(type_id)
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(checklist.resource_path, destination)


def _cell_pixels(ws, coordinate: str) -> tuple[int, int]:
    row, column = coordinate_to_tuple(coordinate)
    width = ws.column_dimensions[get_column_letter(column)].width or 13.0
    height = ws.row_dimensions[row].height or 15.0
    return max(1, int(width * 7 + 5)), max(1, int(height * 96 / 72))


def _add_centered_image(ws, coordinate: str, image_path: Path) -> None:
    cell_width, cell_height = _cell_pixels(ws, coordinate)
    with image_path.open("rb") as handle:
        from PIL import Image as PILImage
        with PILImage.open(handle) as source:
            width, height = source.size
    # Fill each equal-size photo cell as far as possible while keeping the
    # original aspect ratio.  Small source images are deliberately enlarged.
    scale = min((cell_width * 0.92) / max(width, 1), (cell_height * 0.92) / max(height, 1))
    image = ExcelImage(str(image_path))
    image.width = max(1, int(width * scale))
    image.height = max(1, int(height * scale))
    row, column = coordinate_to_tuple(coordinate)
    marker = AnchorMarker(col=column - 1, row=row - 1, colOff=pixels_to_EMU(max(0, int((cell_width - image.width) / 2))), rowOff=pixels_to_EMU(max(0, int((cell_height - image.height) / 2))))
    image.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(cx=pixels_to_EMU(image.width), cy=pixels_to_EMU(image.height)))
    ws.add_image(image)


def export_workbook(type_id: str, selections: dict[str, Path], output_path: Path) -> dict:
    checklist = get_checklist_type(type_id)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(checklist.resource_path, output_path)
    workbook = load_workbook(output_path)
    ws = workbook["Photo report Tool"]
    inserted = 0
    missing: list[str] = []
    for slot in checklist.slots:
        image_path = selections.get(slot.item_key)
        if not image_path or not image_path.is_file():
            missing.append(slot.item_key)
            continue
        _add_centered_image(ws, slot.image_cell, image_path)
        inserted += 1
    workbook.save(output_path)
    return {"output_path": str(output_path), "inserted_count": inserted, "missing_count": len(missing), "missing_item_keys": missing}


def type_payload() -> list[dict]:
    return [{"id": item.id, "label": item.label, "version": item.version, "slot_count": len(item.slots), "sha256": item.sha256} for item in checklist_types()]
