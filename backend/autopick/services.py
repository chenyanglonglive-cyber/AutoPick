from __future__ import annotations

import json
import re
import shutil
import sqlite3
import threading
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from backend.autopick.config import Settings
from backend.autopick.db import connect, initialize_global, initialize_project
from backend.autopick.images import (
    IMAGE_SUFFIXES,
    PREPROCESS_VERSION,
    estimated_visual_tokens,
    flags_json,
    hamming_distance,
    inspect_image,
    make_derivative,
    sha256_file,
)
from backend.autopick.history import extract_history_media
from backend.autopick.report_layout import best_checklist_match, normalize_label
from backend.autopick.local_ocr import LocalOcrClient
from backend.autopick.qwen import QwenEmbeddingClient
from backend.autopick.excel_checklist import copy_template, get_checklist_type, export_workbook


DEFAULT_MATCH_WEIGHTS = {"semantic": 0.62, "quality": 0.10, "ocr": 0.28}
AUTO_CONFIRM_THRESHOLD = 0.35
AUTO_CONFIRM_MIN_SEMANTIC = 0.18


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


class ProjectNotFoundError(FileNotFoundError):
    pass


class JobCancelled(RuntimeError):
    pass


class ProjectService:
    def __init__(self, settings: Settings):
        self.settings = settings
        settings.data_root.mkdir(parents=True, exist_ok=True)
        settings.projects_root.mkdir(parents=True, exist_ok=True)
        initialize_global(settings.global_db)
        self._jobs: dict[str, threading.Thread] = {}

    def project_root(self, project_id: str) -> Path:
        if not project_id or any(value in project_id for value in ("/", "\\", "..")):
            raise ProjectNotFoundError("Invalid project identifier")
        root = self.settings.projects_root / project_id
        database = root / "project.sqlite"
        if not database.exists():
            raise ProjectNotFoundError(f"Project {project_id} was not found")
        # Schema additions must also reach projects created by earlier versions.
        initialize_project(database)
        with connect(database) as db:
            project = db.execute("SELECT checklist_type_id FROM project LIMIT 1").fetchone()
        checklist_type_id = project["checklist_type_id"] if project and project["checklist_type_id"] else "quality_v1"
        with connect(database) as db:
            slot_state = db.execute(
                "SELECT COUNT(*) AS count, SUM(CASE WHEN item_key IS NULL THEN 1 ELSE 0 END) AS legacy FROM checklist_slots WHERE checklist_type_id=?",
                (checklist_type_id,),
            ).fetchone()
        expected_slot_count = len(get_checklist_type(checklist_type_id).slots)
        if int(slot_state["count"] or 0) != expected_slot_count or int(slot_state["legacy"] or 0) > 0:
            self._migrate_fixed_root(root, database, checklist_type_id)
        return root

    @staticmethod
    def _insert_checklist_slots(db: sqlite3.Connection, checklist_type_id: str) -> None:
        checklist = get_checklist_type(checklist_type_id)
        for slot in checklist.slots:
            db.execute(
                "INSERT INTO checklist_slots(id, label, section, ordinal, bookmark, item_key, checklist_type_id, sheet_name, label_cell, image_cell, created_at) VALUES (?, ?, ?, ?, NULL, ?, ?, ?, ?, ?, ?)",
                (str(uuid.uuid4()), slot.label, slot.section, slot.ordinal, slot.item_key, checklist_type_id, slot.sheet_name, slot.label_cell, slot.image_cell, utcnow()),
            )

    def _ensure_project_checklist(self, root: Path, database: Path, checklist_type_id: str) -> None:
        """Add one requested checklist without touching another checklist's evidence."""
        checklist = get_checklist_type(checklist_type_id)
        with connect(database) as db:
            count = int(db.execute(
                "SELECT COUNT(*) FROM checklist_slots WHERE checklist_type_id=?", (checklist_type_id,)
            ).fetchone()[0])
            if count == len(checklist.slots):
                return
            if count:
                raise ValueError(f"清单 {checklist_type_id} 的字段不完整，请先完成项目迁移")
            copy_template(checklist_type_id, root / "templates" / f"{checklist_type_id}.xlsx")
            self._insert_checklist_slots(db, checklist_type_id)

    def _resolve_project_checklist(self, project_id: str, requested_type_id: str | None = None) -> tuple[Path, str]:
        root = self.project_root(project_id)
        database = root / "project.sqlite"
        with connect(database) as db:
            project = db.execute("SELECT checklist_type_id FROM project LIMIT 1").fetchone()
        checklist_type_id = requested_type_id or project["checklist_type_id"]
        get_checklist_type(checklist_type_id)
        self._ensure_project_checklist(root, database, checklist_type_id)
        return root, checklist_type_id

    def _migrate_fixed_root(self, root: Path, database: Path, checklist_type_id: str) -> None:
        checklist = get_checklist_type(checklist_type_id)
        destination = root / "templates" / f"{checklist_type_id}.xlsx"
        copy_template(checklist_type_id, destination)
        with connect(database) as metadata_db:
            project_metadata = metadata_db.execute("SELECT id, factory_name, created_at FROM project LIMIT 1").fetchone()
        if project_metadata:
            (root / "project-info.json").write_text(json.dumps({
                "id": project_metadata["id"], "factory_name": project_metadata["factory_name"],
                "checklist_type_id": checklist_type_id, "created_at": project_metadata["created_at"],
                "data_layout": "originals/<photo_id>__<source_filename>; derived/embedding/<photo_id>__<source_filename>.jpg; project.sqlite",
            }, ensure_ascii=False, indent=2), encoding="utf-8")
        with connect(database) as db:
            for table in ("candidates", "confirmations", "rejections", "preference_feedback", "pending_aliases", "report_slot_manifest", "report_slot_selections", "report_slot_candidates", "report_slots", "report_manifest", "report_runs", "history_matches"):
                db.execute(f"DELETE FROM {table}")
            db.execute("DELETE FROM checklist_slots")
            self._insert_checklist_slots(db, checklist_type_id)
            db.execute("UPDATE project SET template_relative_path=?, checklist_relative_path=?, checklist_type_id=?, schema_version=2", (destination.relative_to(root).as_posix(), destination.relative_to(root).as_posix(), checklist_type_id))

    def db_path(self, project_id: str) -> Path:
        return self.project_root(project_id) / "project.sqlite"

    def create_project(
        self,
        factory_name: str,
        gallery_path: str,
        template_path: str | None = None,
        checklist_path: str | None = None,
        history_report_path: str | None = None,
        checklist_type_id: str = "quality_v1",
    ) -> dict:
        gallery = Path(gallery_path).expanduser().resolve()
        if not gallery.is_dir():
            raise ValueError("图库路径不是有效文件夹")
        checklist = get_checklist_type(checklist_type_id)
        # A compatibility path for pre-Excel callers. The public API no longer
        # exposes this argument, but keeping the file available lets old local
        # projects finish a migration without losing their report metadata.
        legacy_template = Path(template_path).expanduser().resolve() if template_path else None
        if legacy_template and (legacy_template.suffix.lower() not in {".docx", ".docm"} or not legacy_template.is_file()):
            raise ValueError("旧模板必须是存在的 .docx 或 .docm 文件")

        project_id = str(uuid.uuid4())
        root = self.settings.projects_root / project_id
        for child in ("originals", "derived/embedding", "derived/ocr", "templates", "outputs", "manifests"):
            (root / child).mkdir(parents=True, exist_ok=True)
        copied_checklist = root / "templates" / f"{checklist_type_id}.xlsx"
        copy_template(checklist_type_id, copied_checklist)
        copied_template = copied_checklist
        if legacy_template:
            copied_template = root / "templates" / legacy_template.name
            shutil.copy2(legacy_template, copied_template)
        (root / "project-info.json").write_text(json.dumps({
            "id": project_id, "factory_name": factory_name.strip(),
            "checklist_type_id": checklist_type_id, "created_at": utcnow(),
            "data_layout": "originals/<photo_id>__<source_filename>; derived/embedding/<photo_id>__<source_filename>.jpg; project.sqlite",
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        initialize_project(root / "project.sqlite")

        with connect(root / "project.sqlite") as db:
            db.execute(
                "INSERT INTO project(id, factory_name, created_at, status, template_relative_path, checklist_relative_path, embedding_model, embedding_dimension, preprocess_version, checklist_type_id, schema_version) VALUES (?, ?, ?, 'created', ?, ?, ?, ?, ?, ?, 2)",
                (
                    project_id,
                    factory_name.strip(),
                    utcnow(),
                    copied_template.relative_to(root).as_posix(),
                    copied_checklist.relative_to(root).as_posix(),
                    self.settings.embedding_model,
                    self.settings.embedding_dimension,
                    PREPROCESS_VERSION,
                    checklist_type_id,
                ),
            )
            self._insert_checklist_slots(db, checklist_type_id)
        self._sync_catalog([(slot.label, slot.section) for slot in checklist.slots])
        self.import_gallery(project_id, gallery)
        return self.get_project(project_id)

    def _sync_catalog(self, items: Iterable[tuple[str, str | None]]) -> None:
        with connect(self.settings.global_db) as db:
            for label, _ in items:
                db.execute("INSERT OR IGNORE INTO checklist_catalog(label, aliases_json) VALUES (?, '[]')", (label,))

    def import_gallery(self, project_id: str, gallery: Path) -> dict:
        """Add only new image content from a source gallery to one project.

        SHA-256 is the stable content identity. Matching content already has
        local processing records, so it is not copied, vectorized, or OCR'd
        again when the source gallery is updated.
        """
        root = self.project_root(project_id)
        images = sorted(path for path in gallery.rglob("*") if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES)
        if not images:
            raise ValueError("图库中没有可用图片")
        added_count = 0
        skipped_count = 0
        with connect(self.db_path(project_id)) as db:
            known_hashes = {row["sha256"]: row["id"] for row in db.execute("SELECT id, sha256 FROM photos")}
            existing_hashes = set(known_hashes)
            known_dhashes = [(row["id"], row["dhash"]) for row in db.execute("SELECT id, dhash FROM photos")]
            for source in images:
                inspection = inspect_image(source)
                existing = known_hashes.get(inspection.sha256)
                # Only skip content that was already part of this project
                # before this update.  Two files with the same content in one
                # newly imported batch remain visible as separate gallery
                # photos and can share the first photo's vector as before.
                if inspection.sha256 in existing_hashes:
                    skipped_count += 1
                    continue
                photo_id = str(uuid.uuid4())
                target_name = f"{photo_id}__{source.name}"
                target = root / "originals" / target_name
                shutil.copy2(source, target)
                nearby = [known_id for known_id, known_hash in known_dhashes if hamming_distance(inspection.dhash, known_hash) <= 4]
                near_group = nearby[0] if nearby else None
                db.execute(
                    """INSERT INTO photos(id,relative_path,source_filename,sha256,dhash,width,height,byte_size,blur_score,dark_ratio,bright_ratio,quality_score,quality_flags_json,duplicate_of,near_duplicate_group,created_at)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        photo_id,
                        target.relative_to(root).as_posix(),
                        source.name,
                        inspection.sha256,
                        inspection.dhash,
                        inspection.width,
                        inspection.height,
                        inspection.byte_size,
                        inspection.blur_score,
                        inspection.dark_ratio,
                        inspection.bright_ratio,
                        inspection.quality_score,
                        flags_json(inspection.flags),
                        existing,
                        near_group,
                        utcnow(),
                    ),
                )
                known_hashes[inspection.sha256] = photo_id
                known_dhashes.append((photo_id, inspection.dhash))
                added_count += 1
        return {
            "added_count": added_count,
            "skipped_count": skipped_count,
            "message": f"图库已更新：新增 {added_count} 张，已跳过 {skipped_count} 张已有图片。请为新增图片建立向量和本地 OCR。",
        }

    def list_projects(self) -> list[dict]:
        projects = []
        for root in self.settings.projects_root.iterdir():
            if root.is_dir() and (root / "project.sqlite").exists():
                projects.append(self.get_project(root.name))
        return sorted(projects, key=lambda project: project["created_at"], reverse=True)

    def get_project(self, project_id: str) -> dict:
        with connect(self.db_path(project_id)) as db:
            project = dict(db.execute("SELECT * FROM project LIMIT 1").fetchone())
            counts = db.execute(
                """SELECT (SELECT COUNT(*) FROM photos) AS photo_count,
                          (SELECT COUNT(*) FROM embeddings) AS indexed_count,
                          (SELECT COUNT(DISTINCT slot_id) FROM confirmations) AS confirmed_count,
                          (SELECT COUNT(*) FROM checklist_slots) AS slot_count"""
            ).fetchone()
        return {
            **project,
            **dict(counts),
            "template_name": Path(project["template_relative_path"]).name,
        }

    def delete_project(self, project_id: str, factory_name: str) -> dict:
        """Delete one explicitly confirmed project, including its local outputs."""
        root = self.project_root(project_id)
        with connect(self.db_path(project_id)) as db:
            project = db.execute("SELECT factory_name FROM project LIMIT 1").fetchone()
            if not project or project["factory_name"] != factory_name:
                raise ValueError("请输入当前项目名称确认删除项目")
            if db.execute("SELECT 1 FROM jobs WHERE status IN ('queued','running') LIMIT 1").fetchone():
                raise ValueError("当前项目仍有任务运行，暂不能删除")

        target = root.resolve()
        expected_parent = self.settings.projects_root.resolve()
        if target.parent != expected_parent or not target.is_dir():
            raise ValueError("项目目录校验失败，拒绝删除")
        shutil.rmtree(target)
        return {"status": "deleted", "project_id": project_id, "message": f"项目“{factory_name}”已删除"}

    def replace_template(self, project_id: str, template_path: str) -> dict:
        """Store a new report template inside this project and reset only its report mapping."""
        source = Path(template_path).expanduser().resolve()
        if source.suffix.lower() not in {".docx", ".docm"} or not source.is_file():
            raise ValueError("模板必须是存在的 .docx 或 .docm 文件")

        root = self.project_root(project_id)
        destination = root / "templates" / source.name
        if source != destination.resolve():
            shutil.copy2(source, destination)

        with connect(self.db_path(project_id)) as db:
            db.execute(
                "UPDATE project SET template_relative_path=?",
                (destination.relative_to(root).as_posix(),),
            )
            # These rows describe the old template's layout. Historical report
            # runs remain intact, while the next analysis starts from the new one.
            db.execute("DELETE FROM report_slot_candidates")
            db.execute("DELETE FROM report_slot_selections")
            db.execute("DELETE FROM report_slots")

        return {"template_name": destination.name}

    def start_index(self, project_id: str) -> dict:
        root = self.project_root(project_id)
        with connect(self.db_path(project_id)) as db:
            running = db.execute("SELECT * FROM jobs WHERE kind='index' AND status='running' ORDER BY created_at DESC LIMIT 1").fetchone()
            if running:
                return self._job_dict(running, project_id)
            photos = db.execute(
                "SELECT p.* FROM photos p LEFT JOIN embeddings e ON e.photo_id=p.id WHERE e.photo_id IS NULL AND p.duplicate_of IS NULL"
            ).fetchall()
            estimate = sum(estimated_visual_tokens(row["width"], row["height"]) for row in photos)
            if estimate > self.settings.token_budget:
                raise ValueError(f"预计 {estimate} Token，超过当前项目 {self.settings.token_budget} Token 预算")
            job_id = str(uuid.uuid4())
            slot_count = int(db.execute("SELECT COUNT(*) FROM checklist_slots").fetchone()[0])
            db.execute(
                "INSERT INTO jobs VALUES (?, 'index', 'queued', 0, ?, '等待开始', ?, 0, ?, ?)",
                (job_id, len(photos) + slot_count, estimate, utcnow(), utcnow()),
            )
        thread = threading.Thread(target=self._index_worker, args=(project_id, job_id), daemon=True)
        self._jobs[job_id] = thread
        thread.start()
        return self.get_job(project_id, job_id)

    def start_gallery_processing(self, project_id: str) -> dict:
        """Queue one non-blocking pipeline for all unprocessed image vectors, OCR and matching."""
        root = self.project_root(project_id)
        self._copy_exact_duplicate_ocr(project_id)
        with connect(self.db_path(project_id)) as db:
            running = db.execute(
                "SELECT * FROM jobs WHERE kind IN ('index', 'ocr', 'gallery_processing') AND status='running' ORDER BY created_at DESC LIMIT 1"
            ).fetchone()
            if running:
                worker = self._jobs.get(running["id"])
                if worker is not None and worker.is_alive():
                    return self._job_dict(running, project_id)
                db.execute(
                    "UPDATE jobs SET status='interrupted', message=?, updated_at=? WHERE id=?",
                    ("上次图库处理因应用关闭或重启而中断；可重新处理剩余图片", utcnow(), running["id"]),
                )
            vector_rows = db.execute(
                "SELECT p.* FROM photos p LEFT JOIN embeddings e ON e.photo_id=p.id WHERE e.photo_id IS NULL AND p.duplicate_of IS NULL"
            ).fetchall()
            ocr_rows = db.execute(
                "SELECT p.* FROM photos p LEFT JOIN ocr_results o ON o.photo_id=p.id WHERE o.photo_id IS NULL AND p.duplicate_of IS NULL"
            ).fetchall()
            estimate = sum(estimated_visual_tokens(row["width"], row["height"]) for row in vector_rows)
            if estimate > self.settings.token_budget:
                raise ValueError(f"预计 {estimate} Token，超过当前项目 {self.settings.token_budget} Token 预算")
            slot_count = int(db.execute("SELECT COUNT(*) FROM checklist_slots").fetchone()[0])
            job_id = str(uuid.uuid4())
            db.execute(
                "INSERT INTO jobs VALUES (?, 'gallery_processing', 'queued', 0, ?, '等待开始图片向量化', ?, 0, ?, ?)",
                (job_id, len(vector_rows), estimate, utcnow(), utcnow()),
            )
        thread = threading.Thread(target=self._index_worker, args=(project_id, job_id, True), daemon=True)
        self._jobs[job_id] = thread
        thread.start()
        return self.get_job(project_id, job_id)

    def cancel_job(self, project_id: str, job_id: str) -> dict:
        """Request a safe stop between individual image or checklist operations."""
        with connect(self.db_path(project_id)) as db:
            job = db.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
            if not job:
                raise ValueError("任务不存在")
            if job["kind"] != "gallery_processing":
                raise ValueError("只有图片向量化 + OCR 任务可以在此处停止")
            if job["status"] in {"completed", "failed", "cancelled", "interrupted"}:
                return self._job_dict(job, project_id)
            db.execute(
                "UPDATE jobs SET status='cancelling', message='正在停止处理；当前图片完成后将保留已有结果', updated_at=? WHERE id=?",
                (utcnow(), job_id),
            )
        return self.get_job(project_id, job_id)

    def _raise_if_job_cancelled(self, project_id: str, job_id: str) -> None:
        with connect(self.db_path(project_id)) as db:
            row = db.execute("SELECT status FROM jobs WHERE id=?", (job_id,)).fetchone()
        if row and row["status"] == "cancelling":
            raise JobCancelled()

    def _index_worker(self, project_id: str, job_id: str, include_ocr: bool = False) -> None:
        root = self.project_root(project_id)
        with QwenEmbeddingClient(
            self.settings.dashscope_api_key,
            self.settings.embedding_model,
            self.settings.embedding_dimension,
            self.settings.demo_embeddings,
            max_retries=self.settings.qwen_max_retries,
        ) as client:
            try:
                self._raise_if_job_cancelled(project_id, job_id)
                with connect(self.db_path(project_id)) as db:
                    db.execute("UPDATE jobs SET status='running', current=0, message='图片向量化：准备处理待处理图片', updated_at=? WHERE id=?", (utcnow(), job_id))
                    rows = db.execute(
                        "SELECT p.* FROM photos p LEFT JOIN embeddings e ON e.photo_id=p.id WHERE e.photo_id IS NULL AND p.duplicate_of IS NULL ORDER BY p.created_at"
                    ).fetchall()
                for ordinal, row in enumerate(rows, start=1):
                    self._raise_if_job_cancelled(project_id, job_id)
                    original = self.safe_photo_path(project_id, row["id"])
                    derivative = root / "derived" / "embedding" / f"{row['id']}__{row['source_filename']}.jpg"
                    make_derivative(original, derivative, 1024, 82)
                    result = client.embed_image(derivative)
                    with connect(self.db_path(project_id)) as db:
                        db.execute(
                            "INSERT OR REPLACE INTO embeddings VALUES (?, ?, ?, ?, ?, ?, ?)",
                            (
                                row["id"],
                                self.settings.embedding_model,
                                self.settings.embedding_dimension,
                                PREPROCESS_VERSION,
                                result.vector.astype(np.float32).tobytes(),
                                result.input_tokens,
                                utcnow(),
                            ),
                        )
                        db.execute(
                            "UPDATE photos SET embedding_relative_path=?, indexed_at=? WHERE id=?",
                            (derivative.relative_to(root).as_posix(), utcnow(), row["id"]),
                        )
                        db.execute(
                            "INSERT INTO api_usage VALUES (?, ?, ?, 'image_embedding', ?, ?, 'success', NULL, ?)",
                            (str(uuid.uuid4()), job_id, row["id"], self.settings.embedding_model, result.input_tokens, utcnow()),
                        )
                        db.execute(
                            "UPDATE jobs SET current=?, actual_tokens=actual_tokens+?, message=?, updated_at=? WHERE id=?",
                            (ordinal, result.input_tokens, f"图片向量化：{ordinal}/{len(rows)} 张待处理图片", utcnow(), job_id),
                        )
                self._raise_if_job_cancelled(project_id, job_id)
                self._copy_exact_duplicate_vectors(project_id)
                if include_ocr:
                    with connect(self.db_path(project_id)) as db:
                        ocr_rows = db.execute(
                            "SELECT p.* FROM photos p LEFT JOIN ocr_results o ON o.photo_id=p.id WHERE o.photo_id IS NULL AND p.duplicate_of IS NULL ORDER BY p.created_at"
                        ).fetchall()
                        db.execute(
                            "UPDATE jobs SET current=0, total=?, message='本地 OCR：准备识别待处理图片', updated_at=? WHERE id=?",
                            (len(ocr_rows), utcnow(), job_id),
                        )
                    ocr_client = LocalOcrClient()
                    for ordinal, row in enumerate(ocr_rows, start=1):
                        self._raise_if_job_cancelled(project_id, job_id)
                        original = self.safe_photo_path(project_id, row["id"])
                        derivative = root / "derived" / "ocr" / f"{row['id']}__{row['source_filename']}.jpg"
                        make_derivative(original, derivative, 1600, 90)
                        text = ocr_client.extract_text(derivative)
                        with connect(self.db_path(project_id)) as db:
                            db.execute("INSERT OR REPLACE INTO ocr_results VALUES (?, ?, ?)", (row["id"], text, utcnow()))
                            db.execute(
                                "UPDATE jobs SET current=?, message=?, updated_at=? WHERE id=?",
                                (ordinal, f"本地 OCR：{ordinal}/{len(ocr_rows)} 张待处理图片", utcnow(), job_id),
                            )
                    self._raise_if_job_cancelled(project_id, job_id)
                    self._copy_exact_duplicate_ocr(project_id)
                else:
                    processed_image_steps = len(rows)
                with connect(self.db_path(project_id)) as db:
                    slots = db.execute("SELECT id, label, item_key, checklist_type_id FROM checklist_slots ORDER BY ordinal").fetchall()
                    if include_ocr:
                        db.execute(
                            "UPDATE jobs SET current=0, total=?, message='清单匹配：准备重新计算候选照片', updated_at=? WHERE id=?",
                            (len(slots), utcnow(), job_id),
                        )
                    else:
                        db.execute("UPDATE jobs SET message='开始按清单匹配候选照片', updated_at=? WHERE id=?", (utcnow(), job_id))
                for ordinal, slot in enumerate(slots, start=1):
                    self._raise_if_job_cancelled(project_id, job_id)
                    try:
                        with connect(self.settings.global_db) as global_db:
                            aliases = [row["alias"] for row in global_db.execute("SELECT alias FROM checklist_aliases WHERE checklist_type_id=? AND item_key=?", (slot["checklist_type_id"], slot["item_key"] or slot["id"])).fetchall()]
                        query = " ".join([slot["label"], *aliases])
                        self._compute_candidates(project_id, slot["id"], query, 8)
                    except Exception:
                        pass
                    with connect(self.db_path(project_id)) as db:
                        db.execute(
                            "UPDATE jobs SET current=?, total=?, message=?, updated_at=? WHERE id=?",
                            (
                                ordinal if include_ocr else processed_image_steps + ordinal,
                                len(slots) if include_ocr else processed_image_steps + len(slots),
                                f"清单匹配：{ordinal}/{len(slots)} · {slot['label']}" if include_ocr else f"正在匹配清单 {ordinal}/{len(slots)}：{slot['label']}",
                                utcnow(),
                                job_id,
                            ),
                        )
                self._raise_if_job_cancelled(project_id, job_id)
                for checklist_type_id in {slot["checklist_type_id"] for slot in slots}:
                    self._clear_unsuitable_system_confirmations(project_id, checklist_type_id)
                    self._deduplicate_system_confirmations(project_id, checklist_type_id)
                    self.confirm_all_top_candidates(project_id, checklist_type_id)
                with connect(self.db_path(project_id)) as db:
                    completion_message = '图片向量化、OCR 与清单匹配已全部完成' if include_ocr else '图片向量化与清单匹配已全部完成'
                    db.execute("UPDATE jobs SET status='completed', message=?, updated_at=? WHERE id=?", (completion_message, utcnow(), job_id))
                    db.execute("UPDATE project SET status='indexed'")
            except JobCancelled:
                with connect(self.db_path(project_id)) as db:
                    db.execute(
                        "UPDATE jobs SET status='cancelled', message='处理已停止；已完成的向量和 OCR 已保留，可再次点击继续', updated_at=? WHERE id=?",
                        (utcnow(), job_id),
                    )
            except Exception as exc:
                with connect(self.db_path(project_id)) as db:
                    db.execute("UPDATE jobs SET status='failed', message=?, updated_at=? WHERE id=?", (str(exc)[:1000], utcnow(), job_id))

    def _copy_exact_duplicate_vectors(self, project_id: str) -> None:
        with connect(self.db_path(project_id)) as db:
            duplicates = db.execute("SELECT id, duplicate_of FROM photos WHERE duplicate_of IS NOT NULL").fetchall()
            for row in duplicates:
                source = db.execute("SELECT * FROM embeddings WHERE photo_id=?", (row["duplicate_of"],)).fetchone()
                if source:
                    db.execute(
                        "INSERT OR REPLACE INTO embeddings VALUES (?, ?, ?, ?, ?, ?, ?)",
                        (row["id"], source["model"], source["dimension"], source["preprocess_version"], source["vector"], 0, utcnow()),
                    )

    def _copy_exact_duplicate_ocr(self, project_id: str) -> None:
        """Reuse local OCR text for byte-identical gallery photos."""
        with connect(self.db_path(project_id)) as db:
            duplicates = db.execute("SELECT id, duplicate_of FROM photos WHERE duplicate_of IS NOT NULL").fetchall()
            for row in duplicates:
                source = db.execute("SELECT text FROM ocr_results WHERE photo_id=?", (row["duplicate_of"],)).fetchone()
                if source:
                    db.execute(
                        "INSERT OR REPLACE INTO ocr_results VALUES (?, ?, ?)",
                        (row["id"], source["text"], utcnow()),
                    )

    def start_ocr(self, project_id: str) -> dict:
        """Build a local text index only after the user explicitly requests it."""
        root = self.project_root(project_id)
        self._copy_exact_duplicate_ocr(project_id)
        with connect(self.db_path(project_id)) as db:
            running = db.execute("SELECT * FROM jobs WHERE kind='ocr' AND status='running' ORDER BY created_at DESC LIMIT 1").fetchone()
            if running:
                worker = self._jobs.get(running["id"])
                if worker is not None and worker.is_alive():
                    return self._job_dict(running, project_id)
                # Worker threads are process-local.  After an app restart an old
                # daemon thread cannot resume, so leaving this row as "running"
                # would permanently prevent the remaining photos from being read.
                db.execute(
                    "UPDATE jobs SET status='interrupted', message=?, updated_at=? WHERE id=?",
                    ("上次本地 OCR 因应用关闭或重启而中断；已保留已识别图片，可继续处理剩余图片", utcnow(), running["id"]),
                )
            photos = db.execute(
                "SELECT p.* FROM photos p LEFT JOIN ocr_results o ON o.photo_id=p.id WHERE o.photo_id IS NULL AND p.duplicate_of IS NULL"
            ).fetchall()
            job_id = str(uuid.uuid4())
            db.execute(
                "INSERT INTO jobs VALUES (?, 'ocr', 'queued', 0, ?, '等待开始本地文字识别', 0, 0, ?, ?)",
                (job_id, len(photos), utcnow(), utcnow()),
            )
        thread = threading.Thread(target=self._ocr_worker, args=(project_id, job_id), daemon=True)
        self._jobs[job_id] = thread
        thread.start()
        return self.get_job(project_id, job_id)

    def start_match(self, project_id: str, checklist_type_id: str | None = None) -> dict:
        """Re-rank the fixed checklist from existing vectors without re-indexing images."""
        _, checklist_type_id = self._resolve_project_checklist(project_id, checklist_type_id)
        with connect(self.db_path(project_id)) as db:
            running = db.execute("SELECT * FROM jobs WHERE kind='match' AND status='running' ORDER BY created_at DESC LIMIT 1").fetchone()
            if running:
                return self._job_dict(running, project_id)
            if not db.execute("SELECT 1 FROM embeddings LIMIT 1").fetchone():
                raise ValueError("请先建立图片向量，再执行清单匹配")
            total = int(db.execute("SELECT COUNT(*) FROM checklist_slots WHERE checklist_type_id=?", (checklist_type_id,)).fetchone()[0])
            job_id = str(uuid.uuid4())
            db.execute(
                "INSERT INTO jobs VALUES (?, 'match', 'queued', 0, ?, '等待按清单重新匹配', 0, 0, ?, ?)",
                (job_id, total, utcnow(), utcnow()),
            )
        thread = threading.Thread(target=self._match_worker, args=(project_id, job_id, checklist_type_id), daemon=True)
        self._jobs[job_id] = thread
        thread.start()
        return self.get_job(project_id, job_id)

    def _match_worker(self, project_id: str, job_id: str, checklist_type_id: str) -> None:
        try:
            with connect(self.db_path(project_id)) as db:
                db.execute("UPDATE jobs SET status='running', message='正在按清单重新计算候选照片', updated_at=? WHERE id=?", (utcnow(), job_id))
                slots = db.execute(
                    "SELECT id, label, item_key, checklist_type_id FROM checklist_slots WHERE checklist_type_id=? ORDER BY ordinal",
                    (checklist_type_id,),
                ).fetchall()
            for ordinal, slot in enumerate(slots, start=1):
                with connect(self.settings.global_db) as global_db:
                    aliases = [row["alias"] for row in global_db.execute("SELECT alias FROM checklist_aliases WHERE checklist_type_id=? AND item_key=?", (slot["checklist_type_id"], slot["item_key"] or slot["id"])).fetchall()]
                self._compute_candidates(project_id, slot["id"], " ".join([slot["label"], *aliases]), 8)
                with connect(self.db_path(project_id)) as db:
                    db.execute("UPDATE jobs SET current=?, message=?, updated_at=? WHERE id=?", (ordinal, f"正在匹配清单 {ordinal}/{len(slots)}：{slot['label']}", utcnow(), job_id))
            self._clear_unsuitable_system_confirmations(project_id, checklist_type_id)
            self._deduplicate_system_confirmations(project_id, checklist_type_id)
            selected = self.confirm_all_top_candidates(project_id, checklist_type_id)
            with connect(self.db_path(project_id)) as db:
                db.execute("UPDATE jobs SET status='completed', message=?, updated_at=? WHERE id=?", (f"重新匹配完成，已自动选中 {selected['selected_count']} 个最高匹配照片", utcnow(), job_id))
        except Exception as exc:
            with connect(self.db_path(project_id)) as db:
                db.execute("UPDATE jobs SET status='failed', message=?, updated_at=? WHERE id=?", (str(exc)[:1000], utcnow(), job_id))

    def _ocr_worker(self, project_id: str, job_id: str) -> None:
        root = self.project_root(project_id)
        client = LocalOcrClient()
        try:
            with connect(self.db_path(project_id)) as db:
                db.execute("UPDATE jobs SET status='running', message='正在用本地 OCR 建立文字索引', updated_at=? WHERE id=?", (utcnow(), job_id))
                rows = db.execute(
                    "SELECT p.* FROM photos p LEFT JOIN ocr_results o ON o.photo_id=p.id WHERE o.photo_id IS NULL AND p.duplicate_of IS NULL ORDER BY p.created_at"
                ).fetchall()
            for ordinal, row in enumerate(rows, start=1):
                original = self.safe_photo_path(project_id, row["id"])
                derivative = root / "derived" / "ocr" / f"{row['id']}__{row['source_filename']}.jpg"
                make_derivative(original, derivative, 1600, 90)
                text = client.extract_text(derivative)
                with connect(self.db_path(project_id)) as db:
                    db.execute("INSERT OR REPLACE INTO ocr_results VALUES (?, ?, ?)", (row["id"], text, utcnow()))
                    db.execute(
                        "UPDATE jobs SET current=?, message=?, updated_at=? WHERE id=?",
                        (ordinal, f"本地 OCR 已识别 {ordinal}/{len(rows)} 张图片", utcnow(), job_id),
                    )
            self._copy_exact_duplicate_ocr(project_id)
            with connect(self.db_path(project_id)) as db:
                db.execute("UPDATE jobs SET status='completed', message='本地文字索引已完成，可重新匹配清单', updated_at=? WHERE id=?", (utcnow(), job_id))
        except Exception as exc:
            with connect(self.db_path(project_id)) as db:
                db.execute("UPDATE jobs SET status='failed', message=?, updated_at=? WHERE id=?", (str(exc)[:1000], utcnow(), job_id))

    def get_job(self, project_id: str, job_id: str) -> dict:
        with connect(self.db_path(project_id)) as db:
            row = db.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
            if not row:
                raise ValueError("任务不存在")
            return self._job_dict(row, project_id)

    @staticmethod
    def _job_dict(row: sqlite3.Row, project_id: str) -> dict:
        return {"id": row["id"], "project_id": project_id, **dict(row)}

    def list_slots(self, project_id: str, checklist_type_id: str | None = None) -> list[dict]:
        _, checklist_type_id = self._resolve_project_checklist(project_id, checklist_type_id)
        with connect(self.db_path(project_id)) as db:
            rows = db.execute(
                """SELECT s.*,
                          COALESCE((SELECT COUNT(*) FROM candidates c WHERE c.slot_id=s.id),0) candidate_count,
                          (SELECT MAX(c.total_score) FROM candidates c WHERE c.slot_id=s.id) top_score
                   FROM checklist_slots s WHERE s.checklist_type_id=? ORDER BY ordinal""",
                (checklist_type_id,),
            ).fetchall()
            result = []
            for row in rows:
                confirmed = db.execute("SELECT photo_id FROM confirmations WHERE slot_id=?", (row["id"],)).fetchall()
                result.append({**dict(row), "confirmed_photo_ids": [value["photo_id"] for value in confirmed]})
            return result

    def list_gallery(self, project_id: str) -> list[dict]:
        with connect(self.db_path(project_id)) as db:
            rows = db.execute(
                """SELECT p.*, e.photo_id AS embedding_id, o.text AS ocr_text,
                          (SELECT COUNT(*) FROM confirmations c WHERE c.photo_id=p.id) AS usage_count
                   FROM photos p
                   LEFT JOIN embeddings e ON e.photo_id=p.id AND e.model=? AND e.dimension=?
                   LEFT JOIN ocr_results o ON o.photo_id=p.id
                   ORDER BY p.created_at, p.source_filename""",
                (self.settings.embedding_model, self.settings.embedding_dimension),
            ).fetchall()
        return [self._present_gallery_photo(project_id, row) for row in rows]

    @staticmethod
    def _ocr_match_score(query: str, text: str) -> float:
        normalized_query = "".join(query.lower().split())
        normalized_text = "".join(text.lower().split())
        if not normalized_query or not normalized_text:
            return 0.0
        if normalized_query in normalized_text:
            return 1.0
        # Keep IDs (ISO9001, 8.35-39) intact while also allowing meaningful
        # Chinese terms to contribute independently to a partial OCR match.
        tokens = re.findall(r"[a-z0-9][a-z0-9._/-]*|[\u4e00-\u9fff]{2,}", normalized_query)
        tokens = list(dict.fromkeys(token for token in tokens if len(token) > 1))
        if not tokens:
            return 0.0
        hits = sum(token in normalized_text for token in tokens)
        return hits / len(tokens)

    def search_gallery(self, project_id: str, query: str, top_k: int = 120) -> list[dict]:
        self.project_root(project_id)
        with connect(self.db_path(project_id)) as db:
            rows = db.execute(
                """SELECT p.*, e.vector, e.photo_id AS embedding_id, o.text AS ocr_text,
                          (SELECT COUNT(*) FROM confirmations c WHERE c.photo_id=p.id) AS usage_count
                   FROM photos p
                   LEFT JOIN embeddings e ON e.photo_id=p.id AND e.model=? AND e.dimension=?
                   LEFT JOIN ocr_results o ON o.photo_id=p.id""",
                (self.settings.embedding_model, self.settings.embedding_dimension),
            ).fetchall()
        vector_rows = [row for row in rows if row["vector"] is not None]
        semantic: dict[str, float] = {}
        if vector_rows:
            query_vector = self._query_vector(project_id, query)
            matrix = np.vstack([np.frombuffer(row["vector"], dtype=np.float32) for row in vector_rows])
            semantic = {row["id"]: float(score) for row, score in zip(vector_rows, matrix @ query_vector)}
        ranked: list[tuple[float, sqlite3.Row, float, bool]] = []
        for row in rows:
            ocr_score = self._ocr_match_score(query, row["ocr_text"] or "")
            semantic_score = semantic.get(row["id"], -1.0)
            if semantic_score < -0.99 and not ocr_score:
                continue
            total = 0.70 * semantic_score + 0.25 * ocr_score + 0.05 * float(row["quality_score"])
            ranked.append((total, row, semantic_score, bool(ocr_score)))
        ranked.sort(key=lambda item: item[0], reverse=True)
        return [
            self._present_gallery_photo(project_id, row, score=score, semantic_score=semantic_score, ocr_hit=ocr_hit)
            for score, row, semantic_score, ocr_hit in ranked[:top_k]
        ]

    def _present_gallery_photo(
        self, project_id: str, row: sqlite3.Row, score: float | None = None,
        semantic_score: float | None = None, ocr_hit: bool = False,
    ) -> dict:
        text = row["ocr_text"] or ""
        return {
            "photo_id": row["id"],
            "filename": row["source_filename"],
            "preview_url": f"/api/projects/{project_id}/photos/{row['id']}/file",
            "quality_score": round(float(row["quality_score"]), 4),
            "quality_flags": json.loads(row["quality_flags_json"]),
            "embedding_status": "indexed" if row["embedding_id"] else "pending",
            "ocr_status": "done" if row["ocr_text"] is not None else "pending",
            "ocr_text_preview": text[:240],
            "usage_count": int(row["usage_count"]),
            "score": round(float(score), 4) if score is not None else None,
            "semantic_score": round(float(semantic_score), 4) if semantic_score is not None else None,
            "ocr_hit": ocr_hit,
        }

    def candidates(self, project_id: str, slot_id: str, top_k: int = 8) -> list[dict]:
        with connect(self.db_path(project_id)) as db:
            slot = db.execute("SELECT * FROM checklist_slots WHERE id=?", (slot_id,)).fetchone()
            if not slot:
                raise ValueError("清单项目不存在")
            cached = db.execute("SELECT * FROM candidates WHERE slot_id=? ORDER BY rank LIMIT ?", (slot_id, top_k)).fetchall()
        if not cached:
            self._compute_candidates(project_id, slot_id, slot["label"], max(8, top_k))
            with connect(self.db_path(project_id)) as db:
                cached = db.execute("SELECT * FROM candidates WHERE slot_id=? ORDER BY rank LIMIT ?", (slot_id, top_k)).fetchall()
        return self._present_candidates(project_id, cached, slot["checklist_type_id"])

    def _query_vector(self, project_id: str, query: str) -> np.ndarray:
        normalized = " ".join(query.split()).strip().lower()
        with connect(self.db_path(project_id)) as db:
            cached = db.execute("SELECT vector, dimension FROM search_cache WHERE query=?", (normalized,)).fetchone()
            if cached:
                return np.frombuffer(cached["vector"], dtype=np.float32)
        with QwenEmbeddingClient(
            self.settings.dashscope_api_key,
            self.settings.embedding_model,
            self.settings.embedding_dimension,
            self.settings.demo_embeddings,
            max_retries=self.settings.qwen_max_retries,
        ) as client:
            result = client.embed_text(query)
        with connect(self.db_path(project_id)) as db:
            db.execute(
                "INSERT OR REPLACE INTO search_cache VALUES (?, ?, ?, ?, ?)",
                (normalized, result.vector.astype(np.float32).tobytes(), self.settings.embedding_model, self.settings.embedding_dimension, utcnow()),
            )
            db.execute(
                "INSERT INTO api_usage VALUES (?, NULL, NULL, 'text_embedding', ?, ?, 'success', NULL, ?)",
                (str(uuid.uuid4()), self.settings.embedding_model, result.input_tokens, utcnow()),
            )
        return result.vector

    def _vectors_and_photo_rows(self, project_id: str) -> tuple[np.ndarray, list[sqlite3.Row]]:
        with connect(self.db_path(project_id)) as db:
            rows = db.execute(
                """SELECT p.*, e.vector FROM photos p JOIN embeddings e ON e.photo_id=p.id
                   WHERE e.model=? AND e.dimension=?""",
                (self.settings.embedding_model, self.settings.embedding_dimension),
            ).fetchall()
        if not rows:
            raise ValueError("当前项目尚未完成图片向量化")
        matrix = np.vstack([np.frombuffer(row["vector"], dtype=np.float32) for row in rows])
        return matrix, rows

    def _compute_candidates(self, project_id: str, slot_id: str, query: str, top_k: int) -> None:
        query_vector = self._query_vector(project_id, query)
        matrix, rows = self._vectors_and_photo_rows(project_id)
        semantic = matrix @ query_vector
        with connect(self.db_path(project_id)) as db:
            slot = db.execute("SELECT checklist_type_id FROM checklist_slots WHERE id=?", (slot_id,)).fetchone()
            if not slot:
                raise ValueError("清单项目不存在")
        with connect(self.settings.global_db) as global_db:
            profile = global_db.execute("SELECT weights_json FROM preference_profiles WHERE name=?", (slot["checklist_type_id"],)).fetchone()
        weights = {**DEFAULT_MATCH_WEIGHTS, **(json.loads(profile["weights_json"]) if profile else {})}
        semantic_weight = float(weights.get("semantic", DEFAULT_MATCH_WEIGHTS["semantic"]))
        quality_weight = float(weights.get("quality", DEFAULT_MATCH_WEIGHTS["quality"]))
        ocr_weight = float(weights.get("ocr", DEFAULT_MATCH_WEIGHTS["ocr"]))
        with connect(self.db_path(project_id)) as db:
            ocr = {row["photo_id"]: row["text"] for row in db.execute("SELECT photo_id, text FROM ocr_results")}
            rejected_photo_ids = {row["photo_id"] for row in db.execute("SELECT photo_id FROM rejections WHERE slot_id=?", (slot_id,))}
        ranked: list[tuple[float, float, sqlite3.Row]] = []
        for index, row in enumerate(rows):
            if row["id"] in rejected_photo_ids:
                continue
            quality = float(row["quality_score"])
            duplicate_penalty = 0.05 if row["near_duplicate_group"] else 0.0
            ocr_score = self._ocr_match_score(query, ocr.get(row["id"], ""))
            total = (
                semantic_weight * float(semantic[index])
                + ocr_weight * ocr_score
                + quality_weight * quality
                - duplicate_penalty
            )
            ranked.append((total, float(semantic[index]), row))
        ranked.sort(key=lambda item: item[0], reverse=True)
        with connect(self.db_path(project_id)) as db:
            db.execute("DELETE FROM candidates WHERE slot_id=?", (slot_id,))
            for rank, (total, semantic_score, row) in enumerate(ranked[:top_k], start=1):
                db.execute(
                    "INSERT INTO candidates VALUES (?, ?, ?, ?, ?, ?, ?)",
                    (slot_id, row["id"], semantic_score, row["quality_score"], total, rank, utcnow()),
                )

    def search(self, project_id: str, query: str, top_k: int, checklist_type_id: str | None = None) -> list[dict]:
        query_vector = self._query_vector(project_id, query)
        matrix, rows = self._vectors_and_photo_rows(project_id)
        semantic = matrix @ query_vector
        ranked: list[tuple[float, float, sqlite3.Row, bool]] = []
        with connect(self.db_path(project_id)) as db:
            ocr = {row["photo_id"]: row["text"].lower() for row in db.execute("SELECT photo_id, text FROM ocr_results")}
        for index, row in enumerate(rows):
            ocr_score = self._ocr_match_score(query, ocr.get(row["id"], ""))
            score = 0.70 * float(semantic[index]) + 0.25 * ocr_score + 0.05 * float(row["quality_score"])
            ranked.append((score, float(semantic[index]), row, bool(ocr_score)))
        ranked.sort(key=lambda item: item[0], reverse=True)
        synthetic = []
        for rank, (score, semantic_score, row, ocr_hit) in enumerate(ranked[:top_k], start=1):
            synthetic.append({"photo_id": row["id"], "total_score": score, "semantic_score": semantic_score, "quality_score": row["quality_score"], "rank": rank, "ocr_hit": ocr_hit})
        return self._present_candidates(project_id, synthetic, checklist_type_id)

    def _present_candidates(
        self, project_id: str, candidates: Iterable[sqlite3.Row | dict], checklist_type_id: str | None = None,
    ) -> list[dict]:
        result: list[dict] = []
        with connect(self.db_path(project_id)) as db:
            for candidate in candidates:
                photo_id = candidate["photo_id"]
                photo = db.execute("SELECT * FROM photos WHERE id=?", (photo_id,)).fetchone()
                if not photo:
                    continue
                if checklist_type_id:
                    used = db.execute(
                        "SELECT c.slot_id FROM confirmations c JOIN checklist_slots s ON s.id=c.slot_id WHERE c.photo_id=? AND s.checklist_type_id=?",
                        (photo_id, checklist_type_id),
                    ).fetchall()
                else:
                    used = db.execute("SELECT slot_id FROM confirmations WHERE photo_id=?", (photo_id,)).fetchall()
                result.append(
                    {
                        "photo_id": photo_id,
                        "filename": photo["source_filename"],
                        "preview_url": f"/api/projects/{project_id}/photos/{photo_id}/file",
                        "score": round(float(candidate["total_score"]), 4),
                        "semantic_score": round(float(candidate["semantic_score"]), 4),
                        "quality_score": round(float(candidate["quality_score"]), 4),
                        "ocr_hit": bool(candidate.get("ocr_hit", False)) if isinstance(candidate, dict) else False,
                        "used_in_slots": [row["slot_id"] for row in used],
                        "quality_flags": json.loads(photo["quality_flags_json"]),
                    }
                )
        return result

    def confirm_all_top_candidates(self, project_id: str, checklist_type_id: str | None = None) -> dict:
        """Confirm the best available candidates without reusing a photo automatically."""
        _, checklist_type_id = self._resolve_project_checklist(project_id, checklist_type_id)
        with connect(self.db_path(project_id)) as db:
            rows = db.execute(
                """SELECT s.id AS slot_id, c.photo_id, c.total_score, c.rank
                   FROM checklist_slots s
                   JOIN candidates c ON c.slot_id=s.id
                   WHERE s.checklist_type_id=?
                     AND NOT EXISTS (SELECT 1 FROM confirmations x WHERE x.slot_id=s.id)
                     AND NOT EXISTS (SELECT 1 FROM rejections r WHERE r.slot_id=s.id AND r.photo_id=c.photo_id)
                     AND c.total_score>=?
                     AND c.semantic_score>=?
                   ORDER BY c.total_score DESC, c.rank ASC, s.ordinal""",
                (checklist_type_id, AUTO_CONFIRM_THRESHOLD, AUTO_CONFIRM_MIN_SEMANTIC),
            ).fetchall()
            reserved_photo_ids = {
                row["photo_id"] for row in db.execute(
                    "SELECT c.photo_id FROM confirmations c JOIN checklist_slots s ON s.id=c.slot_id WHERE s.checklist_type_id=? AND c.source='system'",
                    (checklist_type_id,),
                )
            }
            pending = int(db.execute(
                "SELECT COUNT(*) FROM checklist_slots s WHERE s.checklist_type_id=? AND NOT EXISTS (SELECT 1 FROM confirmations x WHERE x.slot_id=s.id)",
                (checklist_type_id,),
            ).fetchone()[0])
            total = int(db.execute("SELECT COUNT(*) FROM checklist_slots WHERE checklist_type_id=?", (checklist_type_id,)).fetchone()[0])
        selected_rows = []
        selected_slot_ids: set[str] = set()
        for row in rows:
            if row["slot_id"] in selected_slot_ids or row["photo_id"] in reserved_photo_ids:
                continue
            selected_rows.append(row)
            selected_slot_ids.add(row["slot_id"])
            reserved_photo_ids.add(row["photo_id"])
        for row in selected_rows:
            self.confirm(project_id, row["slot_id"], [row["photo_id"]], source="system")
        return {
            "selected_count": len(selected_rows),
            "skipped_confirmed_count": total - pending,
            "unmatched_count": max(0, pending - len(selected_rows)),
            "message": f"已为 {len(selected_rows)} 个未确认清单项选中不重复照片",
        }

    def _deduplicate_system_confirmations(self, project_id: str, checklist_type_id: str) -> int:
        """Remove stale duplicate auto-selections while preserving all manual choices."""
        with connect(self.db_path(project_id)) as db:
            rows = db.execute(
                """SELECT c.slot_id, c.photo_id, c.confirmed_at, COALESCE(k.total_score, -999.0) AS total_score
                   FROM confirmations c
                   JOIN checklist_slots s ON s.id=c.slot_id
                   LEFT JOIN candidates k ON k.slot_id=c.slot_id AND k.photo_id=c.photo_id
                   WHERE s.checklist_type_id=? AND c.source='system'
                   ORDER BY c.photo_id, total_score DESC, c.confirmed_at ASC""",
                (checklist_type_id,),
            ).fetchall()
            retained_photo_ids: set[str] = set()
            duplicate_slot_ids: list[str] = []
            for row in rows:
                if row["photo_id"] in retained_photo_ids:
                    duplicate_slot_ids.append(row["slot_id"])
                else:
                    retained_photo_ids.add(row["photo_id"])
            for slot_id in duplicate_slot_ids:
                db.execute("DELETE FROM confirmations WHERE slot_id=? AND source='system'", (slot_id,))
        return len(duplicate_slot_ids)

    def _clear_unsuitable_system_confirmations(self, project_id: str, checklist_type_id: str) -> int:
        """Leave a field blank when its current system choice no longer meets the threshold."""
        with connect(self.db_path(project_id)) as db:
            cleared = db.execute(
                """DELETE FROM confirmations
                   WHERE source='system'
                     AND slot_id IN (SELECT id FROM checklist_slots WHERE checklist_type_id=?)
                     AND NOT EXISTS (
                         SELECT 1 FROM candidates c
                         WHERE c.slot_id=confirmations.slot_id
                           AND c.photo_id=confirmations.photo_id
                           AND c.total_score>=?
                           AND c.semantic_score>=?
                     )""",
                (checklist_type_id, AUTO_CONFIRM_THRESHOLD, AUTO_CONFIRM_MIN_SEMANTIC),
            ).rowcount
        return int(cleared)

    def confirm(self, project_id: str, slot_id: str, photo_ids: list[str], source: str = "manual", search_query: str | None = None) -> None:
        root = self.project_root(project_id)
        with connect(self.db_path(project_id)) as db:
            if not db.execute("SELECT 1 FROM checklist_slots WHERE id=?", (slot_id,)).fetchone():
                raise ValueError("清单项目不存在")
            for photo_id in photo_ids:
                photo = db.execute("SELECT relative_path, sha256 FROM photos WHERE id=?", (photo_id,)).fetchone()
                if not photo:
                    raise ValueError("照片不属于当前项目")
                path = self._safe_relative(root, photo["relative_path"])
                if not path.exists() or sha256_file(path) != photo["sha256"]:
                    raise ValueError("照片文件校验失败，无法确认")
            system = db.execute("SELECT photo_id FROM confirmations WHERE slot_id=? ORDER BY confirmed_at LIMIT 1", (slot_id,)).fetchone()
            db.execute("DELETE FROM confirmations WHERE slot_id=?", (slot_id,))
            for photo_id in photo_ids:
                db.execute(
                    "INSERT INTO confirmations(slot_id, photo_id, confirmed_at, source, search_query, system_photo_id) VALUES (?, ?, ?, ?, ?, ?)",
                    (slot_id, photo_id, utcnow(), source, search_query, system["photo_id"] if system else photo_id),
                )
                if source == "manual" and system and system["photo_id"] != photo_id:
                    checklist_type = db.execute("SELECT checklist_type_id FROM checklist_slots WHERE id=?", (slot_id,)).fetchone()["checklist_type_id"]
                    exists = db.execute("SELECT 1 FROM preference_feedback WHERE slot_id=? AND applied=0", (slot_id,)).fetchone()
                    if not exists:
                        db.execute(
                            "INSERT INTO preference_feedback(id, checklist_type_id, slot_id, selected_photo_id, system_photo_id, created_at) VALUES (?, ?, ?, ?, ?, ?)",
                            (str(uuid.uuid4()), checklist_type, slot_id, photo_id, system["photo_id"], utcnow()),
                        )
            if source == "manual" and search_query and search_query.strip():
                slot = db.execute("SELECT item_key, checklist_type_id FROM checklist_slots WHERE id=?", (slot_id,)).fetchone()
                db.execute(
                    "INSERT INTO pending_aliases(id, checklist_type_id, item_key, alias, created_at) VALUES (?, ?, ?, ?, ?)",
                    (str(uuid.uuid4()), slot["checklist_type_id"], slot["item_key"] or slot_id, search_query.strip(), utcnow()),
                )

    def reject(self, project_id: str, slot_id: str, photo_id: str, reason: str | None) -> None:
        with connect(self.db_path(project_id)) as db:
            db.execute(
                "INSERT OR REPLACE INTO rejections VALUES (?, ?, ?, ?)",
                (slot_id, photo_id, reason, utcnow()),
            )

    def import_excel_feedback(self, project_id: str, feedback_path: str) -> dict:
        """Record retained and deleted images against their original export batch."""
        self.project_root(project_id)
        source = Path(feedback_path).expanduser().resolve()
        if source.suffix.lower() != ".xlsx" or not source.is_file():
            raise ValueError("用户反馈必须是存在的 .xlsx 清单文件")
        if source.stat().st_size > 100 * 1024 * 1024:
            raise ValueError("用户反馈文件超过 100 MB，无法导入")
        source_sha256 = sha256_file(source)

        try:
            workbook = load_workbook(source, read_only=False, data_only=False)
        except Exception as exc:
            raise ValueError(f"无法读取用户反馈 Excel：{exc}") from exc
        image_cells: set[tuple[str, str]] = set()
        for worksheet in workbook.worksheets:
            for image in worksheet._images:
                anchor = getattr(image, "anchor", None)
                origin = getattr(anchor, "_from", None)
                if origin is not None:
                    image_cells.add((worksheet.title, f"{get_column_letter(origin.col + 1)}{origin.row + 1}"))

        with connect(self.db_path(project_id)) as db:
            if db.execute("SELECT 1 FROM jobs WHERE status IN ('queued','running') LIMIT 1").fetchone():
                raise ValueError("当前项目仍有任务运行，暂不能导入用户反馈")
            if db.execute("SELECT 1 FROM feedback_imports WHERE source_sha256=?", (source_sha256,)).fetchone():
                raise ValueError("这份用户反馈已经导入过")
            feedback_cells = image_cells
            batches = db.execute(
                "SELECT id, output_relative_path, checklist_type_id FROM export_batches ORDER BY created_at DESC"
            ).fetchall()
            # Users commonly append “修改” to the exported filename. Keep the
            # timestamp as a stable link to the exact original export when it
            # is still present; otherwise use the image-cell overlap fallback.
            stamp_match = re.search(r"\d{8}-\d{6}", source.stem)
            if stamp_match:
                stamp = stamp_match.group(0)
                timestamp_batches = [
                    batch
                    for batch in batches
                    if stamp in Path(batch["output_relative_path"]).stem
                ]
                if timestamp_batches:
                    batches = timestamp_batches
            best_batch: str | None = None
            best_output_name = ""
            best_items: list[sqlite3.Row] = []
            best_overlap = -1
            for batch in batches:
                items = db.execute(
                    """SELECT i.*, s.sheet_name, s.image_cell
                       FROM export_batch_items i JOIN checklist_slots s ON s.id=i.slot_id
                       WHERE i.export_id=?""",
                    (batch["id"],),
                ).fetchall()
                item_cells = {(item["sheet_name"], item["image_cell"]) for item in items}
                overlap = len(feedback_cells.intersection(item_cells))
                if feedback_cells and overlap == 0:
                    continue
                if overlap > best_overlap:
                    best_batch, best_output_name, best_items, best_overlap = batch["id"], Path(batch["output_relative_path"]).name, items, overlap
            if not best_batch or not best_items:
                raise ValueError("未找到可对应的原始导出批次。请使用升级后重新导出的清单收集用户反馈")
            required_sheets = {item["sheet_name"] for item in best_items}
            if not required_sheets.issubset(set(workbook.sheetnames)):
                raise ValueError("该 Excel 缺少对应清单工作表，请确认导入的是本系统导出的清单")

            import_id = str(uuid.uuid4())
            db.execute(
                "INSERT INTO feedback_imports(id, source_sha256, source_filename, export_id, imported_at) VALUES (?, ?, ?, ?, ?)",
                (import_id, source_sha256, source.name, best_batch, utcnow()),
            )
            accepted_count = 0
            rejected_count = 0
            for item in best_items:
                cell = (item["sheet_name"], item["image_cell"])
                retained = cell in feedback_cells
                outcome = "accepted" if retained else "rejected"
                db.execute(
                    """INSERT INTO feedback_events(id, import_id, export_id, slot_id, system_photo_id, selected_photo_id, outcome, source, created_at)
                       VALUES (?, ?, ?, ?, ?, ?, ?, 'excel_feedback', ?)""",
                    (
                        str(uuid.uuid4()), import_id, best_batch, item["slot_id"], item["system_photo_id"],
                        item["selected_photo_id"] if retained else None, outcome, utcnow(),
                    ),
                )
                if retained:
                    accepted_count += 1
                    continue
                rejected_count += 1
                db.execute(
                    "INSERT OR REPLACE INTO rejections VALUES (?, ?, ?, ?)",
                    (item["slot_id"], item["selected_photo_id"], "用户反馈 Excel 中已删除图片", utcnow()),
                )
                # If a later manual correction already changed this field, keep it.
                db.execute(
                    "DELETE FROM confirmations WHERE slot_id=? AND photo_id=?",
                    (item["slot_id"], item["selected_photo_id"]),
                )
                db.execute("DELETE FROM candidates WHERE slot_id=?", (item["slot_id"],))

        return {
            "status": "imported",
            "export_id": best_batch,
            "baseline_export_name": best_output_name,
            "accepted_count": accepted_count,
            "rejected_count": rejected_count,
            "message": f"已导入用户反馈：确认正确 {accepted_count} 项，标记错误 {rejected_count} 项；重新匹配时将排除错误图片。",
        }

    def set_mapping(self, project_id: str, slot_id: str, bookmark: str) -> None:
        with connect(self.db_path(project_id)) as db:
            changed = db.execute("UPDATE checklist_slots SET bookmark=? WHERE id=?", (bookmark, slot_id)).rowcount
            if changed != 1:
                raise ValueError("清单项目不存在")

    def export_preflight(self, project_id: str, checklist_type_id: str | None = None) -> dict:
        _, checklist_type_id = self._resolve_project_checklist(project_id, checklist_type_id)
        with connect(self.db_path(project_id)) as db:
            project = db.execute("SELECT factory_name, checklist_type_id FROM project LIMIT 1").fetchone()
            total = int(db.execute("SELECT COUNT(*) FROM checklist_slots WHERE checklist_type_id=?", (checklist_type_id,)).fetchone()[0])
            selected = int(db.execute(
                "SELECT COUNT(DISTINCT c.slot_id) FROM confirmations c JOIN checklist_slots s ON s.id=c.slot_id WHERE s.checklist_type_id=?",
                (checklist_type_id,),
            ).fetchone()[0])
            overrides = int(db.execute(
                "SELECT COUNT(*) FROM confirmations c JOIN checklist_slots s ON s.id=c.slot_id WHERE s.checklist_type_id=? AND c.source='manual' AND c.system_photo_id != c.photo_id",
                (checklist_type_id,),
            ).fetchone()[0])
            pending_aliases = int(db.execute(
                "SELECT COUNT(*) FROM pending_aliases WHERE checklist_type_id=? AND confirmed_at IS NULL",
                (checklist_type_id,),
            ).fetchone()[0])
            jobs = db.execute("SELECT COUNT(*) FROM jobs WHERE status IN ('queued','running')").fetchone()[0]
            feedback_records = int(db.execute(
                "SELECT COUNT(*) FROM feedback_events e JOIN checklist_slots s ON s.id=e.slot_id WHERE s.checklist_type_id=?",
                (checklist_type_id,),
            ).fetchone()[0])
        training_feedback = 0
        training_history = 0
        for database in self.settings.projects_root.glob("*/project.sqlite"):
            try:
                with connect(database) as training_db:
                    training_feedback += int(training_db.execute(
                        "SELECT COUNT(*) FROM feedback_events e JOIN checklist_slots s ON s.id=e.slot_id WHERE s.checklist_type_id=?",
                        (checklist_type_id,),
                    ).fetchone()[0])
                    training_history += int(training_db.execute(
                        "SELECT COUNT(*) FROM history_matches h JOIN checklist_slots s ON s.id=h.matched_slot_id WHERE h.photo_id IS NOT NULL AND h.matched_slot_id IS NOT NULL AND s.checklist_type_id=?",
                        (checklist_type_id,),
                    ).fetchone()[0])
            except sqlite3.Error:
                # A legacy or partially migrated project must not hide readiness
                # for the healthy projects that still contain usable feedback.
                continue
        training_data_count = training_feedback + training_history
        training_data_threshold = self.settings.training_feedback_threshold
        training_ready = training_data_count >= training_data_threshold
        training_remaining = max(0, training_data_threshold - training_data_count)
        training_progress_percent = min(100, round(training_data_count / training_data_threshold * 100))
        return {
            "ready": True,
            "total_slots": total,
            "selected_slots": selected,
            "missing_slots": max(0, total - selected),
            "active_overrides": overrides,
            "pending_aliases": pending_aliases,
            "feedback_records": feedback_records,
            "training_data_count": training_data_count,
            "training_feedback_count": training_feedback,
            "training_history_count": training_history,
            "training_data_threshold": training_data_threshold,
            "training_remaining": training_remaining,
            "training_progress_percent": training_progress_percent,
            "training_ready": training_ready,
            "active_jobs": int(jobs),
            "message": f"已积累 {feedback_records} 条有效反馈；本次主动换图 {overrides} 项；发现 {pending_aliases} 个待确认别名。",
            "allow_partial": True,
            "factory_name": project["factory_name"],
            "checklist_type_id": checklist_type_id,
        }

    def export_excel(self, project_id: str, allow_partial: bool = True, checklist_type_id: str | None = None) -> dict:
        root, checklist_type_id = self._resolve_project_checklist(project_id, checklist_type_id)
        with connect(self.db_path(project_id)) as db:
            project = db.execute("SELECT factory_name, checklist_type_id FROM project LIMIT 1").fetchone()
            rows = db.execute(
                """SELECT s.id AS slot_id, s.item_key, p.relative_path, c.photo_id AS selected_photo_id,
                          c.system_photo_id, k.semantic_score, k.quality_score, k.total_score, k.rank
                   FROM checklist_slots s JOIN confirmations c ON c.slot_id=s.id JOIN photos p ON p.id=c.photo_id
                   LEFT JOIN candidates k ON k.slot_id=s.id AND k.photo_id=c.photo_id
                   WHERE s.checklist_type_id=? ORDER BY s.ordinal""",
                (checklist_type_id,),
            ).fetchall()
            total = int(db.execute("SELECT COUNT(*) FROM checklist_slots WHERE checklist_type_id=?", (checklist_type_id,)).fetchone()[0])
        if not allow_partial and len(rows) < total:
            raise ValueError(f"尚有 {total - len(rows)} 个清单项未选图")
        selections = {row["item_key"]: self._safe_relative(root, row["relative_path"]) for row in rows}
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        output = root / "outputs" / f"{project['factory_name']}__{checklist_type_id}__{stamp}.xlsx"
        result = export_workbook(checklist_type_id, selections, output)
        export_id = str(uuid.uuid4())
        with connect(self.db_path(project_id)) as db:
            pending_feedback = db.execute("SELECT * FROM preference_feedback WHERE checklist_type_id=? AND applied=0", (checklist_type_id,)).fetchall()
            batch_overrides = len(pending_feedback)
            batch_aliases = int(db.execute("SELECT COUNT(*) FROM pending_aliases WHERE checklist_type_id=? AND confirmed_at IS NULL", (checklist_type_id,)).fetchone()[0])
            db.execute("INSERT INTO export_batches(id, output_relative_path, checklist_type_id, created_at) VALUES (?, ?, ?, ?)", (export_id, output.relative_to(root).as_posix(), checklist_type_id, utcnow()))
            for row in rows:
                db.execute(
                    """INSERT INTO export_batch_items(export_id, slot_id, selected_photo_id, system_photo_id, semantic_score, quality_score, total_score, rank)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                    (export_id, row["slot_id"], row["selected_photo_id"], row["system_photo_id"], row["semantic_score"], row["quality_score"], row["total_score"], row["rank"]),
                )
            for feedback in pending_feedback:
                db.execute(
                    """INSERT INTO feedback_events(id, export_id, slot_id, system_photo_id, selected_photo_id, outcome, source, created_at)
                       VALUES (?, ?, ?, ?, ?, 'replaced', 'manual', ?)""",
                    (str(uuid.uuid4()), export_id, feedback["slot_id"], feedback["system_photo_id"], feedback["selected_photo_id"], utcnow()),
                )
            db.execute("UPDATE preference_feedback SET applied=1, export_id=? WHERE checklist_type_id=? AND applied=0", (export_id, checklist_type_id))
            db.execute("UPDATE pending_aliases SET confirmed_at=? WHERE checklist_type_id=? AND confirmed_at IS NULL", (utcnow(), checklist_type_id))
            aliases = db.execute("SELECT checklist_type_id, item_key, alias FROM pending_aliases WHERE checklist_type_id=? AND confirmed_at IS NOT NULL AND alias != ''", (checklist_type_id,)).fetchall()
            feedback_count = int(db.execute(
                "SELECT COUNT(*) FROM feedback_events e JOIN checklist_slots s ON s.id=e.slot_id WHERE s.checklist_type_id=?",
                (checklist_type_id,),
            ).fetchone()[0])
        with connect(self.settings.global_db) as global_db:
            for alias in aliases:
                global_db.execute(
                    "INSERT OR IGNORE INTO checklist_aliases(checklist_type_id, item_key, alias, source) VALUES (?, ?, ?, 'search')",
                    (alias["checklist_type_id"], alias["item_key"], alias["alias"]),
                )
        result.update(self.export_preflight(project_id, checklist_type_id))
        result.update({"export_id": export_id, "output_path": str(output), "feedback_records": feedback_count, "message": f"本次主动换图 {batch_overrides} 项；已积累有效反馈 {feedback_count} 条；发现 {batch_aliases} 个待确认别名；缺图 {result.get('missing_count', 0)} 项。"})
        return result

    def clear_gallery(self, project_id: str, factory_name: str) -> dict:
        root = self.project_root(project_id)
        with connect(self.db_path(project_id)) as db:
            project = db.execute("SELECT factory_name FROM project LIMIT 1").fetchone()
            if not project or project["factory_name"] != factory_name:
                raise ValueError("请输入当前项目名称确认清空图库")
            running = db.execute("SELECT 1 FROM jobs WHERE status IN ('queued','running') LIMIT 1").fetchone()
            if running:
                raise ValueError("向量化或文字识别任务运行中，暂不能清空图库")
            photos = [row["relative_path"] for row in db.execute("SELECT relative_path FROM photos")]
            derived = [row["embedding_relative_path"] for row in db.execute("SELECT embedding_relative_path FROM photos WHERE embedding_relative_path IS NOT NULL")]
            for table in ("candidates", "confirmations", "rejections", "ocr_results", "embeddings", "photos", "search_cache", "preference_feedback", "pending_aliases"):
                db.execute(f"DELETE FROM {table}")
            db.execute("UPDATE project SET status='created'")
        for relative in photos + derived:
            if relative:
                path = self._safe_relative(root, relative)
                if path.is_file():
                    path.unlink()
        for folder in (root / "originals", root / "derived" / "embedding", root / "derived" / "ocr"):
            for path in folder.glob("*"):
                if path.is_file():
                    path.unlink()
        return {"status": "cleared", "deleted_photos": len(photos), "message": "当前项目图库、向量和 OCR 已清空，外部源图库与 Excel 输出未修改。"}

    def convert_project(self, project_id: str, checklist_type_id: str = "quality_v1") -> dict:
        """Rebind an older Word-based project to the fixed Excel checklist.

        Photo files, embeddings, OCR and API usage are intentionally untouched;
        only checklist/matching/report metadata is rebuilt.
        """
        root = self.project_root(project_id)
        checklist = get_checklist_type(checklist_type_id)
        destination = root / "templates" / f"{checklist_type_id}.xlsx"
        copy_template(checklist_type_id, destination)
        with connect(self.db_path(project_id)) as db:
            db.execute("DELETE FROM candidates")
            db.execute("DELETE FROM confirmations")
            db.execute("DELETE FROM rejections")
            db.execute("DELETE FROM preference_feedback")
            db.execute("DELETE FROM pending_aliases")
            for table in ("report_slot_manifest", "report_slot_selections", "report_slot_candidates", "report_slots", "report_manifest", "report_runs", "history_matches"):
                db.execute(f"DELETE FROM {table}")
            db.execute("DELETE FROM checklist_slots")
            for slot in checklist.slots:
                db.execute(
                    "INSERT INTO checklist_slots(id, label, section, ordinal, bookmark, item_key, checklist_type_id, sheet_name, label_cell, image_cell, created_at) VALUES (?, ?, ?, ?, NULL, ?, ?, ?, ?, ?, ?)",
                    (str(uuid.uuid4()), slot.label, slot.section, slot.ordinal, slot.item_key, checklist_type_id, slot.sheet_name, slot.label_cell, slot.image_cell, utcnow()),
                )
            db.execute("UPDATE project SET template_relative_path=?, checklist_relative_path=?, checklist_type_id=?, schema_version=2, status=CASE WHEN EXISTS(SELECT 1 FROM embeddings) THEN 'indexed' ELSE 'created' END", (destination.relative_to(root).as_posix(), destination.relative_to(root).as_posix(), checklist_type_id))
        return self.get_project(project_id)

    def template_bookmarks(self, project_id: str) -> list[str]:
        root = self.project_root(project_id)
        with connect(self.db_path(project_id)) as db:
            project = db.execute("SELECT template_relative_path FROM project LIMIT 1").fetchone()
        template = self._safe_relative(root, project["template_relative_path"])
        try:
            with zipfile.ZipFile(template) as package:
                document = ET.fromstring(package.read("word/document.xml"))
        except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
            raise ValueError("模板不是可读取的Word文件") from exc
        namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        names = [node.attrib.get(f"{{{namespace}}}name", "") for node in document.iter(f"{{{namespace}}}bookmarkStart")]
        return sorted(name for name in names if name and not name.startswith("_"))

    def learn_from_report(self, project_id: str, report_path: Path) -> dict:
        """Associate embedded historical-report media with this project's photos only."""
        root = self.project_root(project_id)
        report = report_path.resolve()
        if root.resolve() not in report.parents:
            raise ValueError("历史报告必须先复制到当前项目目录")
        media = extract_history_media(report)
        with connect(self.db_path(project_id)) as db:
            photos = db.execute("SELECT id, sha256, dhash FROM photos").fetchall()
            slots = db.execute(
                "SELECT id, label, item_key, checklist_type_id FROM checklist_slots"
            ).fetchall()
            db.execute("DELETE FROM history_matches WHERE report_relative_path=?", (report.relative_to(root).as_posix(),))
            matched = 0
            for item in media:
                exact = next((photo for photo in photos if photo["sha256"] == item.sha256), None)
                nearest = None
                distance = None
                if exact:
                    nearest, distance = exact, 0
                elif item.dhash:
                    options = sorted(
                        ((hamming_distance(item.dhash, photo["dhash"]), photo) for photo in photos),
                        key=lambda value: value[0],
                    )
                    if options and options[0][0] <= 8:
                        distance, nearest = options[0]
                slot_id, mapping_score = best_checklist_match(
                    item.context_text, [(slot["id"], slot["label"]) for slot in slots]
                )
                confidence = 1.0 if distance == 0 else (max(0.0, 1 - (distance or 99) / 16) if distance is not None else 0.0)
                if nearest:
                    matched += 1
                db.execute(
                    "INSERT INTO history_matches VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (
                        str(uuid.uuid4()),
                        report.relative_to(root).as_posix(),
                        item.media_name,
                        nearest["id"] if nearest else None,
                        slot_id,
                        item.context_text,
                        distance,
                        confidence,
                        utcnow(),
                    ),
                )
                # Only generic field language is copied to the global catalog.
                # Photo IDs, paths, hashes and supplier-specific report data stay in this project DB.
                if slot_id and mapping_score >= 0.82 and item.context_text:
                    label = next(slot["label"] for slot in slots if slot["id"] == slot_id)
                    with connect(self.settings.global_db) as global_db:
                        catalog = global_db.execute(
                            "SELECT aliases_json FROM checklist_catalog WHERE label=?", (label,)
                        ).fetchone()
                        aliases = json.loads(catalog["aliases_json"]) if catalog else []
                        alias = normalize_label(item.context_text)
                        if alias and len(alias) <= 100 and alias not in aliases:
                            aliases = (aliases + [alias])[-20:]
                            global_db.execute(
                                "INSERT INTO checklist_catalog(label, aliases_json) VALUES (?, ?) "
                                "ON CONFLICT(label) DO UPDATE SET aliases_json=excluded.aliases_json",
                                (label, json.dumps(aliases, ensure_ascii=False)),
                            )
                        if alias and len(alias) <= 100:
                            matched_slot = next(slot for slot in slots if slot["id"] == slot_id)
                            if matched_slot["item_key"]:
                                global_db.execute(
                                    "INSERT OR IGNORE INTO checklist_aliases"
                                    "(checklist_type_id, item_key, alias, source) VALUES (?, ?, ?, 'history')",
                                    (
                                        matched_slot["checklist_type_id"],
                                        matched_slot["item_key"],
                                        alias,
                                    ),
                                )
        return {"media_count": len(media), "matched_count": matched}

    @staticmethod
    def _context_slot(context: str, slots: Iterable[sqlite3.Row]) -> str | None:
        normalized = " ".join(context.lower().split())
        matches = [slot for slot in slots if slot["label"].lower() in normalized]
        return matches[0]["id"] if matches else None

    def safe_photo_path(self, project_id: str, photo_id: str) -> Path:
        root = self.project_root(project_id)
        with connect(self.db_path(project_id)) as db:
            photo = db.execute("SELECT relative_path FROM photos WHERE id=?", (photo_id,)).fetchone()
            if not photo:
                raise ValueError("照片不属于当前项目")
        return self._safe_relative(root, photo["relative_path"])

    @staticmethod
    def _safe_relative(root: Path, relative_path: str) -> Path:
        candidate = (root / relative_path).resolve()
        if root.resolve() not in candidate.parents:
            raise ValueError("检测到跨项目文件路径，操作已阻止")
        return candidate
