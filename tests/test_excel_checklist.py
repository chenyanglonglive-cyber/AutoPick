import shutil
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from backend.autopick.config import Settings
from backend.autopick.db import connect
from backend.autopick.excel_checklist import checklist_types
from backend.autopick.services import ProjectService


def _image(path: Path, color=(40, 100, 180)) -> None:
    Image.new("RGB", (120, 80), color).save(path)


def test_quality_registry_has_185_unique_slots_and_skips_markers() -> None:
    checklist = next(item for item in checklist_types() if item.id == "quality_v1")
    assert checklist.id == "quality_v1"
    assert checklist.label == "Quality Checklist"
    assert len(checklist.slots) == 185
    assert all(slot.label.lower() != "onsite:" for slot in checklist.slots)
    assert len({slot.item_key for slot in checklist.slots}) == 185
    assert len({slot.label.lower() for slot in checklist.slots}) == 185
    # 验证重复项如“电池测试证书”仅保留 1 个，多余的已清除
    battery_slots = [slot for slot in checklist.slots if slot.label == "电池测试证书"]
    assert len(battery_slots) == 1
    assert all(slot.image_cell.endswith(str(int(slot.label_cell[1:]) - 1)) for slot in checklist.slots)


def test_social_audit_registry_has_26_reserved_photo_slots() -> None:
    checklist = next(item for item in checklist_types() if item.id == "social_audit_v1")
    assert checklist.label == "Social Audit Checklist"
    assert len(checklist.slots) == 26
    assert {slot.sheet_name for slot in checklist.slots} == {"Sheet1"}
    assert len({slot.item_key for slot in checklist.slots}) == 26
    assert len({slot.label.lower() for slot in checklist.slots}) == 26


def test_excel_export_preserves_sheets_and_inserts_image(tmp_path: Path) -> None:
    gallery = tmp_path / "gallery"
    gallery.mkdir()
    _image(gallery / "evidence.jpg")
    service = ProjectService(Settings(tmp_path / "data", None, True, "token"))
    project = service.create_project("Factory", str(gallery))
    slot = service.list_slots(project["id"])[0]
    photo_id = service.list_gallery(project["id"])[0]["photo_id"]
    service.confirm(project["id"], slot["id"], [photo_id])
    result = service.export_excel(project["id"])
    workbook = load_workbook(result["output_path"])
    assert workbook.sheetnames == ["Photo report Tool", "Sheet2"]
    assert len(workbook["Photo report Tool"]._images) == 1
    assert result["inserted_count"] == 1
    assert result["missing_count"] == 184


def test_social_audit_export_inserts_into_reserved_sheet1_cell(tmp_path: Path) -> None:
    gallery = tmp_path / "gallery"
    gallery.mkdir()
    _image(gallery / "evidence.jpg")
    service = ProjectService(Settings(tmp_path / "data", None, True, "token"))
    project = service.create_project("Factory", str(gallery), checklist_type_id="social_audit_v1")
    slot = service.list_slots(project["id"])[0]
    photo_id = service.list_gallery(project["id"])[0]["photo_id"]
    service.confirm(project["id"], slot["id"], [photo_id])
    result = service.export_excel(project["id"])
    workbook = load_workbook(result["output_path"])
    assert workbook.sheetnames == ["Sheet1", "Sheet2", "Sheet3"]
    assert len(workbook["Sheet1"]._images) == 1
    assert result["inserted_count"] == 1
    assert result["missing_count"] == 25


def test_project_can_switch_checklists_without_sharing_confirmations_or_exports(tmp_path: Path) -> None:
    gallery = tmp_path / "gallery"
    gallery.mkdir()
    _image(gallery / "evidence.jpg")
    service = ProjectService(Settings(tmp_path / "data", None, True, "token"))
    project = service.create_project("Factory", str(gallery), checklist_type_id="quality_v1")
    photo_id = service.list_gallery(project["id"])[0]["photo_id"]

    quality_slot = service.list_slots(project["id"], "quality_v1")[0]
    social_slot = service.list_slots(project["id"], "social_audit_v1")[0]
    service.confirm(project["id"], quality_slot["id"], [photo_id])
    service.confirm(project["id"], social_slot["id"], [photo_id])

    quality_preflight = service.export_preflight(project["id"], "quality_v1")
    social_preflight = service.export_preflight(project["id"], "social_audit_v1")
    assert (quality_preflight["total_slots"], quality_preflight["selected_slots"]) == (185, 1)
    assert (social_preflight["total_slots"], social_preflight["selected_slots"]) == (26, 1)

    quality_export = service.export_excel(project["id"], checklist_type_id="quality_v1")
    social_export = service.export_excel(project["id"], checklist_type_id="social_audit_v1")
    assert quality_export["missing_count"] == 184
    assert social_export["missing_count"] == 25


def test_rematch_removes_duplicate_automatic_selections_within_one_checklist(tmp_path: Path) -> None:
    gallery = tmp_path / "gallery"
    gallery.mkdir()
    _image(gallery / "evidence.jpg")
    service = ProjectService(Settings(tmp_path / "data", None, True, "token"))
    project = service.create_project("Factory", str(gallery))
    slots = service.list_slots(project["id"], "quality_v1")[:2]
    photo_id = service.list_gallery(project["id"])[0]["photo_id"]
    with connect(service.db_path(project["id"])) as db:
        for slot in slots:
            db.execute(
                "INSERT INTO candidates VALUES (?, ?, ?, ?, ?, ?, ?)",
                (slot["id"], photo_id, 0.9, 1.0, 0.6, 1, "now"),
            )
            db.execute(
                "INSERT INTO confirmations(slot_id, photo_id, confirmed_at, source, system_photo_id) VALUES (?, ?, ?, 'system', ?)",
                (slot["id"], photo_id, "now", photo_id),
            )

    assert service._deduplicate_system_confirmations(project["id"], "quality_v1") == 1
    assert sum(bool(slot["confirmed_photo_ids"]) for slot in service.list_slots(project["id"], "quality_v1")[:2]) == 1


def test_automatic_selection_leaves_low_confidence_field_blank(tmp_path: Path) -> None:
    gallery = tmp_path / "gallery"
    gallery.mkdir()
    _image(gallery / "evidence.jpg")
    service = ProjectService(Settings(tmp_path / "data", None, True, "token"))
    project = service.create_project("Factory", str(gallery))
    slot = service.list_slots(project["id"], "quality_v1")[0]
    photo_id = service.list_gallery(project["id"])[0]["photo_id"]
    with connect(service.db_path(project["id"])) as db:
        db.execute(
            "INSERT INTO candidates VALUES (?, ?, ?, ?, ?, ?, ?)",
            (slot["id"], photo_id, 0.17, 1.0, 0.34, 1, "now"),
        )

    result = service.confirm_all_top_candidates(project["id"], "quality_v1")
    assert result["selected_count"] == 0
    assert service.list_slots(project["id"], "quality_v1")[0]["confirmed_photo_ids"] == []


def test_clear_gallery_keeps_export_and_external_source(tmp_path: Path) -> None:
    gallery = tmp_path / "gallery"
    gallery.mkdir()
    source = gallery / "evidence.jpg"
    _image(source)
    service = ProjectService(Settings(tmp_path / "data", None, True, "token"))
    project = service.create_project("Factory", str(gallery))
    result = service.clear_gallery(project["id"], "Factory")
    assert result["deleted_photos"] == 1
    assert source.is_file()
    assert service.get_project(project["id"])["photo_count"] == 0


def test_skip_slot_leaves_slot_empty_on_export_and_avoids_auto_selection(tmp_path: Path) -> None:
    gallery = tmp_path / "gallery"
    gallery.mkdir()
    _image(gallery / "evidence.jpg")
    service = ProjectService(Settings(tmp_path / "data", None, True, "token"))
    project = service.create_project("Factory", str(gallery))
    slot = service.list_slots(project["id"], "quality_v1")[0]
    photo_id = service.list_gallery(project["id"])[0]["photo_id"]

    # 先确认选图
    service.confirm(project["id"], slot["id"], [photo_id])
    assert service.list_slots(project["id"], "quality_v1")[0]["confirmed_photo_ids"] == [photo_id]

    # 跳过该项
    service.skip_slot(project["id"], slot["id"])
    slots = service.list_slots(project["id"], "quality_v1")
    assert slots[0]["is_skipped"] is True
    assert slots[0]["confirmed_photo_ids"] == []

    # 导出 Excel 检查缺图数量包含跳过留空项
    result = service.export_excel(project["id"])
    assert result["inserted_count"] == 0
    assert result["missing_count"] == 185

    # 验证自动选最佳不会为跳过的槽位填图
    with connect(service.db_path(project["id"])) as db:
        db.execute(
            "INSERT INTO candidates VALUES (?, ?, ?, ?, ?, ?, ?)",
            (slot["id"], photo_id, 0.9, 1.0, 0.9, 1, "now"),
        )
    auto_res = service.confirm_all_top_candidates(project["id"], "quality_v1")
    assert auto_res["selected_count"] == 0
    assert service.list_slots(project["id"], "quality_v1")[0]["confirmed_photo_ids"] == []

    # 取消跳过或人工选图
    service.unskip_slot(project["id"], slot["id"])
    assert service.list_slots(project["id"], "quality_v1")[0]["is_skipped"] is False


def test_import_photos_multi_selection_and_deduplication(tmp_path: Path) -> None:
    gallery = tmp_path / "gallery"
    gallery.mkdir()
    _image(gallery / "existing_a.jpg")
    service = ProjectService(Settings(tmp_path / "data", None, True, "token"))
    project = service.create_project("Factory", str(gallery))

    # Verify initial photo count is 1
    photos = service.list_gallery(project["id"])
    assert len(photos) == 1
    assert photos[0]["filename"] == "existing_a.jpg"

    # Prepare external image files:
    # 1. new_b.jpg (distinct new image)
    # 2. new_b_dup.jpg (same content as new_b.jpg - intra-batch duplicate)
    # 3. existing_a_dup.jpg (same content as existing_a.jpg - cross-batch duplicate)
    ext_dir = tmp_path / "external"
    ext_dir.mkdir()
    new_b = ext_dir / "new_b.jpg"
    new_b_dup = ext_dir / "new_b_dup.jpg"
    existing_a_dup = ext_dir / "existing_a_dup.jpg"

    Image.new("RGB", (600, 600), (120, 150, 180)).save(new_b, "JPEG")
    shutil.copy2(new_b, new_b_dup)
    shutil.copy2(gallery / "existing_a.jpg", existing_a_dup)

    # Call import_photos with all three files
    res = service.import_photos(project["id"], [new_b, new_b_dup, existing_a_dup])
    assert res["added_count"] == 1
    assert res["skipped_count"] == 2
    assert "成功新增 1 张图片，已自动跳过 2 张重复照片" in res["message"]

    # Verify project now has exactly 2 photos (existing_a and new_b)
    current_photos = service.list_gallery(project["id"])
    assert len(current_photos) == 2
    filenames = {p["filename"] for p in current_photos}
    assert filenames == {"existing_a.jpg", "new_b.jpg"}

    # Re-importing the same files again skips all 3
    res_again = service.import_photos(project["id"], [new_b, new_b_dup, existing_a_dup])
    assert res_again["added_count"] == 0
    assert res_again["skipped_count"] == 3
    assert "已自动跳过 3 张已有重复照片" in res_again["message"]
